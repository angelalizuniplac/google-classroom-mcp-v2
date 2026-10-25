"""Converte bytes de um arquivo em FileItem(s) legiveis: codigo, texto, PDF, DOCX,
notebook, CSV, XLSX, imagem e ZIP."""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import PurePosixPath

from .config import (CODE_EXTENSIONS, MAX_IMAGE_BYTES, MAX_ZIP_MEMBERS, MAX_ZIP_UNCOMPRESSED,
                     SKIP_DIRS, SKIP_FILENAMES, SPECIAL_FILENAMES, TEXT_EXTENSIONS)
from .models import FileItem
from .notebook import notebook_warnings, parse_notebook
from .safety import scan_text
from .textutil import shrink_image

IMAGE_EXTS = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
              ".gif": "image/gif", ".webp": "image/webp"}
NOTEBOOK_MIMES = {"application/vnd.google.colaboratory", "application/x-ipynb+json"}
TABLE_SAMPLE_ROWS = 20
PROSE_EXTS = {".txt", ".md", ".rst", ".tex"}


def ext_of(name: str) -> str:
    return PurePosixPath(name.replace("\\", "/")).suffix.lower()


def base_of(name: str) -> str:
    return PurePosixPath(name.replace("\\", "/")).name.lower()


def is_readable_name(name: str) -> bool:
    """Decide se um nome de arquivo vale a pena ser lido (repositorios e zips)."""
    parts = PurePosixPath(name.replace("\\", "/")).parts
    if any(p in SKIP_DIRS or p == "__MACOSX" for p in parts[:-1]):
        return False
    base = base_of(name)
    if base in SKIP_FILENAMES or base.endswith((".min.js", ".min.css", ".map")):
        return False
    if base.startswith("._"):
        return False
    e = ext_of(name)
    return (e in CODE_EXTENSIONS or e in TEXT_EXTENSIONS or e in IMAGE_EXTS
            or e in (".pdf", ".docx", ".xlsx", ".zip") or base in SPECIAL_FILENAMES)


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _is_binary(data: bytes) -> bool:
    return b"\x00" in data[:4096]


# --------------------------------------------------------------------------
# Leitores por tipo
# --------------------------------------------------------------------------

def read_pdf(name: str, data: bytes) -> FileItem:
    item = FileItem(name=name, kind="pdf", size=len(data), plag_kind="prose")
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                pages.append("[Pagina %d]\n%s" % (i, t))
        item.meta["paginas"] = len(reader.pages)
        item.text = "\n\n".join(pages)
        item.plag_text = item.text
        if not item.text:
            item.warnings.append("PDF sem texto extraivel (provavelmente digitalizado). Paginas renderizadas como imagem, se possivel.")
            item.meta["scanned"] = True
            item.images = _render_pdf_pages(data, max_pages=4)
            item.plag_kind = ""
    except Exception as e:
        item.error = "falha ao ler PDF: %s" % e
    item.warnings += scan_text(item.text)
    return item


def _render_pdf_pages(data: bytes, max_pages: int = 4) -> list:
    """Renderiza paginas como PNG se pypdfium2 estiver instalado (opcional)."""
    try:
        import pypdfium2 as pdfium
    except Exception:
        return []
    out = []
    try:
        pdf = pdfium.PdfDocument(data)
        for i in range(min(len(pdf), max_pages)):
            img = pdf[i].render(scale=1.3).to_pil()
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            small = shrink_image(buf.getvalue(), "image/png")
            if small:
                out.append(small)
    except Exception:
        pass
    return out


def read_docx(name: str, data: bytes) -> FileItem:
    item = FileItem(name=name, kind="docx", size=len(data), plag_kind="prose")
    try:
        import docx
        from docx.shared import RGBColor
        doc = docx.Document(io.BytesIO(data))
        parts, hidden = [], 0
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text)
            for r in p.runs:
                try:
                    white = r.font.color is not None and r.font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF)
                    if r.text.strip() and (white or r.font.hidden):
                        hidden += len(r.text)
                except Exception:
                    pass
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)
        item.text = "\n".join(parts)
        item.plag_text = item.text
        item.meta["imagens"] = len(doc.inline_shapes)
        if hidden:
            item.warnings.append("%d caracteres de texto oculto (cor branca ou marcado como oculto) no DOCX" % hidden)
        if not item.text:
            item.warnings.append("DOCX sem texto (pode conter apenas imagens)")
    except Exception as e:
        item.error = ("falha ao ler DOCX: %s. Se for .doc antigo, peca para salvar como .docx ou PDF." % e)
    item.warnings += scan_text(item.text)
    return item


def read_notebook(name: str, data: bytes) -> FileItem:
    item = FileItem(name=name, kind="notebook", size=len(data))
    try:
        nb = parse_notebook(data)
        item.text = nb["text"]
        item.plag_text = nb["code"]
        item.plag_kind = "code"
        item.meta = {"stats": nb["stats"], "code_cells": nb["code_cells"]}
        item.warnings += notebook_warnings(nb["stats"])
    except Exception as e:
        item.error = "notebook invalido ou corrompido: %s" % e
    item.warnings += scan_text(item.text)
    return item


def read_table(name: str, data: bytes, delimiter: str | None = None) -> FileItem:
    item = FileItem(name=name, kind="table", size=len(data))
    text = decode_text(data)
    if delimiter is None:
        delimiter = "\t" if ext_of(name) == ".tsv" else (";" if text[:2000].count(";") > text[:2000].count(",") else ",")
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    n = len(rows)
    ncols = max((len(r) for r in rows[:50]), default=0)
    sample = rows[: TABLE_SAMPLE_ROWS + 1]
    lines = ["[tabela: %d linhas x %d colunas; amostra das primeiras %d linhas]" % (n, ncols, len(sample))]
    lines += [delimiter.join(r) for r in sample]
    item.text = "\n".join(lines)
    item.meta = {"linhas": n, "colunas": ncols, "cabecalho": rows[0] if rows else []}
    return item


def read_xlsx(name: str, data: bytes) -> FileItem:
    item = FileItem(name=name, kind="table", size=len(data))
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > TABLE_SAMPLE_ROWS:
                    break
                rows.append(["" if v is None else str(v) for v in row])
            out.append("[planilha '%s': %s linhas x %s colunas; amostra]\n%s" % (
                ws.title, ws.max_row, ws.max_column, "\n".join(",".join(r) for r in rows)))
        item.text = "\n\n".join(out)
    except Exception as e:
        item.error = "falha ao ler XLSX: %s" % e
    return item


def read_image(name: str, data: bytes, mime: str = "") -> FileItem:
    item = FileItem(name=name, kind="image", size=len(data))
    m = mime if mime.startswith("image/") else IMAGE_EXTS.get(ext_of(name), "image/png")
    if len(data) > MAX_IMAGE_BYTES * 5:
        item.warnings.append("imagem grande demais (%d bytes), omitida" % len(data))
        return item
    small = shrink_image(data, m)
    if small:
        item.images = [small]
        item.meta["reduzida"] = True
    else:
        item.warnings.append("imagem nao pode ser reduzida, omitida para poupar tokens")
    return item


def read_text_like(name: str, data: bytes) -> FileItem:
    e = ext_of(name)
    item = FileItem(name=name, size=len(data))
    if _is_binary(data):
        item.kind = "other"
        item.warnings.append("arquivo binario, nao lido")
        return item
    item.text = decode_text(data)
    if e in CODE_EXTENSIONS:
        item.kind, item.plag_kind = "code", "code"
        item.plag_text = item.text
    elif e in PROSE_EXTS:
        item.kind, item.plag_kind = "text", "prose"
        item.plag_text = item.text
    else:
        item.kind = "text"
    item.warnings += scan_text(item.text)
    return item


def read_zip(name: str, data: bytes, depth: int = 0) -> list[FileItem]:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except Exception as e:
        return [FileItem(name=name, kind="other", size=len(data), error="zip invalido: %s" % e)]
    infos = [i for i in zf.infolist() if not i.is_dir()]
    if len(infos) > MAX_ZIP_MEMBERS * 5 or sum(i.file_size for i in infos) > MAX_ZIP_UNCOMPRESSED * 3:
        return [FileItem(name=name, kind="other", size=len(data),
                         error="zip grande demais (possivel zip bomb), nao extraido")]
    items: list[FileItem] = []
    total = 0
    skipped = 0
    for info in infos:
        if not is_readable_name(info.filename):
            skipped += 1
            continue
        if len(items) >= MAX_ZIP_MEMBERS or total + info.file_size > MAX_ZIP_UNCOMPRESSED:
            skipped += 1
            continue
        if ext_of(info.filename) == ".zip" and depth >= 1:
            skipped += 1
            continue
        total += info.file_size
        try:
            member = zf.read(info)
        except Exception as e:
            items.append(FileItem(name="%s::%s" % (name, info.filename), kind="other", error=str(e)))
            continue
        for it in read_bytes("%s::%s" % (name, info.filename), member, source="zip", depth=depth + 1):
            it.source = "zip"
            items.append(it)
    if skipped:
        items.append(FileItem(name=name, kind="other",
                              warnings=["%d arquivo(s) do zip ignorados (binarios, dependencias ou limite)" % skipped]))
    return items


def read_bytes(name: str, data: bytes, mime: str = "", source: str = "", depth: int = 0) -> list[FileItem]:
    """Ponto de entrada: escolhe o leitor pelo nome/mime e devolve FileItems."""
    e = ext_of(name)
    mime = (mime or "").lower()
    if e == ".zip" or mime in ("application/zip", "application/x-zip-compressed"):
        items = read_zip(name, data, depth)
    elif e == ".ipynb" or mime in NOTEBOOK_MIMES:
        items = [read_notebook(name, data)]
    elif e == ".pdf" or mime == "application/pdf":
        items = [read_pdf(name, data)]
    elif e == ".docx" or "wordprocessingml" in mime:
        items = [read_docx(name, data)]
    elif e == ".doc" or mime == "application/msword":
        items = [FileItem(name=name, kind="other", size=len(data),
                          error=".doc antigo nao e suportado. Peca para salvar como .docx ou PDF.")]
    elif e == ".xlsx" or "spreadsheetml" in mime:
        items = [read_xlsx(name, data)]
    elif e in (".csv", ".tsv") or mime in ("text/csv", "text/tab-separated-values"):
        items = [read_table(name, data)]
    elif e in IMAGE_EXTS or mime.startswith("image/"):
        items = [read_image(name, data, mime)]
    elif e in CODE_EXTENSIONS or e in TEXT_EXTENSIONS or base_of(name) in SPECIAL_FILENAMES or mime.startswith("text/") \
            or mime in ("application/json", "application/xml", "application/javascript"):
        items = [read_text_like(name, data)]
    else:
        items = [FileItem(name=name, kind="other", size=len(data),
                          warnings=["tipo nao suportado: %s" % (mime or e or "desconhecido")])]
    for it in items:
        if not it.source:
            it.source = source
    return items
